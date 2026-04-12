from docx import Document
import os

text = (
    "Describe the idea and the system requirements:\n\n"
    "This project is a lightweight data-quality and validation application that allows users to upload tabular datasets (CSV/XLSX) and quickly assess and improve their data without requiring machine learning or embedded AI components. "
    "Through a simple web-based interface, users can inspect dataset summaries (row/column counts, data types, and sample values), detect missing values and duplicates, identify outliers using standard statistical methods, and validate numeric constraints such as ranges, allowed values, and uniqueness. "
    "The system prioritizes clarity and minimal configuration: users select a file, view an automatic diagnostic report, apply optional cleaning operations or filters interactively, and export a cleaned dataset or a concise summary report with visualizations.\n\n"

    "Key features include automatic detection of missingness and duplicate rows, outlier detection using interquartile range and z-score methods with configurable thresholds, numeric constraint checks (min/max, step sizes, categorical allowed values), and generation of summary reports with histograms, boxplots, and missing-value breakdowns. "
    "The UI provides sampling and preview capabilities to inspect records before and after cleaning, as well as simple metadata annotations for columns. Users can download cleaned CSVs and an HTML/PDF report showing the diagnostics and visual evidence.\n\n"

    "Functional requirements:\n"
    "- Accept CSV and Excel (XLSX) uploads; support common encodings and separator options.\n"
    "- Produce dataset overview: number of rows/columns, inferred data types, example values, and counts of missing values per column.\n"
    "- Detect and report duplicate rows and duplicates on user-selected key columns.\n"
    "- Detect outliers per numeric column (IQR and z-score) and allow threshold adjustments.\n"
    "- Validate numeric constraints (range checks, uniqueness, allowed categorical values).\n"
    "- Generate and export a summary report (HTML or PDF) with visualizations and an option to export a cleaned CSV.\n\n"

    "Non-functional requirements:\n"
    "- Performance: reasonably responsive for files up to 100k rows; surface progress feedback for larger files.\n"
    "- Portability: runnable locally and deployable in a container; minimal external dependencies.\n"
    "- Privacy: uploaded files are processed locally and not transmitted to external services.\n"
    "- Reliability: core behaviors covered by automated unit tests (pytest).\n"
    "- Usability: intuitive web UI (Streamlit) with clear error messages and an accessible README for reproducibility.\n\n"

    "Technology stack: Python 3.11+, pandas and numpy for data processing, Streamlit for the UI (or FastAPI + lightweight frontend if separated), matplotlib/seaborn for charts, and pytest for testing. Development was assisted using AI coding tools (GitHub Copilot, Claude Code, Cursor) for productivity, but the delivered application includes no AI models or external AI services as part of runtime."
)

os.makedirs('deliverables', exist_ok=True)

doc = Document()
doc.add_heading('Idea and System Requirements', level=1)
for para in text.split('\n\n'):
    doc.add_paragraph(para)

out_path = os.path.join('deliverables', 'Describe_Idea_and_Requirements.docx')
doc.save(out_path)
print('Saved:', out_path)
